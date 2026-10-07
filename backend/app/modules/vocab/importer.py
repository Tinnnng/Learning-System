"""讀取單字表（CSV／Excel），欄位名稱中英文都接受。"""

import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from app.modules.vocab.schemas import WordIn

COLUMN_ALIASES = {
    "term": {"term", "word", "單字", "英文", "vocabulary", "英文單字", "生字"},
    "meaning": {"meaning", "中文", "意思", "解釋", "definition", "翻譯", "中文意思", "中文解釋", "字義", "意義", "中譯"},
    "part_of_speech": {"part_of_speech", "pos", "詞性"},
    "example": {"example", "例句", "sentence"},
    "example_translation": {"example_translation", "例句翻譯", "例句中文"},
    "tags": {"tags", "tag", "標籤", "分類"},
}


class ImportFormatError(ValueError):
    pass


def _decode(data: bytes) -> str:
    # 台灣版 Excel 另存 CSV 預設是 Big5（cp950），UTF-8 失敗時改用 cp950
    for encoding in ("utf-8-sig", "cp950"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFormatError("無法判斷檔案編碼，請另存成 UTF-8 的 CSV")


def _rows_from_bytes(filename: str, data: bytes) -> list[list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return [row for row in csv.reader(io.StringIO(_decode(data)))]
    if suffix == ".xlsx":
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    raise ImportFormatError(f"不支援 {suffix}，請使用 .csv 或 .xlsx")


def _column_map(header: list[str]) -> dict[str, int]:
    mapping = {}
    for index, name in enumerate(header):
        key = name.strip().lower()
        for field, aliases in COLUMN_ALIASES.items():
            if key in aliases and field not in mapping:
                mapping[field] = index
    missing = [field for field in ("term", "meaning") if field not in mapping]
    if missing:
        raise ImportFormatError(f"第一列必須有欄位名稱，至少要有 term（單字）與 meaning（中文），找到：{header}")
    return mapping


def parse_word_table(filename: str, data: bytes) -> tuple[list[WordIn], list[str]]:
    """回傳 (單字, 錯誤訊息)。有問題的列略過並記錄，不會讓整個檔案失敗。"""
    rows = [row for row in _rows_from_bytes(filename, data) if any(cell.strip() for cell in row)]
    if not rows:
        raise ImportFormatError("檔案是空的")
    columns = _column_map(rows[0])
    words, errors = [], []
    for line_number, row in enumerate(rows[1:], start=2):
        def cell(field: str) -> str | None:
            index = columns.get(field)
            value = row[index].strip() if index is not None and index < len(row) else ""
            return value or None

        tags = [tag.strip() for tag in re.split(r"[|,、;；]", cell("tags") or "") if tag.strip()]
        try:
            words.append(
                WordIn(
                    term=cell("term") or "",
                    meaning=cell("meaning") or "",
                    part_of_speech=cell("part_of_speech"),
                    example=cell("example"),
                    example_translation=cell("example_translation"),
                    tags=tags,
                )
            )
        except ValidationError:
            errors.append(f"第 {line_number} 列缺少單字或中文意思，已略過")
    return words, errors
