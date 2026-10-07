"""抽出上課教材的文字，並標示投影片編號或頁碼，給 Claude Code 出題時引用出處。"""

import subprocess
import sys
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation


def _pptx(path: Path, image_dir: Path | None) -> list[str]:
    sections = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
            if shape.shape_type == 13:  # 圖片
                lines.append(_save_picture(shape, number, image_dir))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("備忘稿：" + slide.notes_slide.notes_text_frame.text.strip())
        sections.append(f"## 投影片 {number}\n\n" + "\n".join(lines))
    return sections


def _save_picture(shape, slide_number: int, image_dir: Path | None) -> str:
    """有 image_dir 就把圖片存出來（掃描成圖片的投影片要看圖才能出題），回傳要寫進 Markdown 的標記。"""
    if image_dir is None:
        return "[圖片]"
    image_dir.mkdir(parents=True, exist_ok=True)
    index = len(list(image_dir.glob(f"slide-{slide_number:02d}-*"))) + 1
    target = image_dir / f"slide-{slide_number:02d}-{index}.{shape.image.ext}"
    target.write_bytes(shape.image.blob)
    return f"[圖片：{target.as_posix()}]"


def _docx(path: Path, image_dir: Path | None) -> list[str]:
    # Word 檔沒有固定頁碼，用標題切段落，出處寫「第 N 段」或標題名稱
    document = DocxDocument(path)
    sections: list[str] = []
    current: list[str] = []
    title = "開頭"
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.lower().startswith(("heading", "標題")):
            if current:
                sections.append(f"## {title}\n\n" + "\n".join(current))
            title, current = text, []
        else:
            current.append(text)
    if current:
        sections.append(f"## {title}\n\n" + "\n".join(current))
    for index, table in enumerate(document.tables, start=1):
        rows = ["| " + " | ".join(cell.text.strip() for cell in row.cells) + " |" for row in table.rows]
        sections.append(f"## 表格 {index}\n\n" + "\n".join(rows))
    return sections


def _pdf(path: Path, image_dir: Path | None) -> list[str]:
    with pymupdf.open(path) as document:
        return [f"## 第 {number} 頁\n\n{page.get_text().strip()}" for number, page in enumerate(document, start=1)]


_EXTRACTORS = {".pptx": _pptx, ".docx": _docx, ".pdf": _pdf}


# 舊版 Office 格式：用本機的 PowerPoint／Word（COM）另存成新格式再抽文字。
# 值是 (COM 程式, 另存格式代碼, 新副檔名)；24 = ppSaveAsOpenXMLPresentation，16 = wdFormatDocumentDefault
_LEGACY = {
    ".ppt": ("PowerPoint.Application", 24, ".pptx"),
    ".doc": ("Word.Application", 16, ".docx"),
}


def convert_legacy(path: Path, out_dir: Path) -> Path:
    """把 .ppt/.doc 另存成 .pptx/.docx（放在 out_dir），回傳新檔路徑。已轉過且比原檔新就直接用。"""
    app, file_format, suffix = _LEGACY[path.suffix.lower()]
    if not path.exists():
        raise ValueError(f"找不到 {path}")
    target = (out_dir / path.with_suffix(suffix).name).resolve()
    if target.exists() and target.stat().st_mtime >= path.stat().st_mtime:
        return target
    if sys.platform != "win32":
        raise ValueError(f"{path.suffix} 需要在 Windows 上用 Office 轉檔，或請先另存成 {suffix}")
    out_dir.mkdir(parents=True, exist_ok=True)
    source = str(path.resolve()).replace("'", "''")
    dest = str(target).replace("'", "''")
    if suffix == ".pptx":
        # Presentations.Open(FileName, ReadOnly, Untitled, WithWindow)
        steps = f"$doc = $app.Presentations.Open('{source}', -1, 0, 0); $doc.SaveAs('{dest}', {file_format})"
    else:
        steps = f"$app.Visible = $false; $doc = $app.Documents.Open('{source}', $false, $true); $doc.SaveAs2('{dest}', {file_format})"
    script = (
        f"$ErrorActionPreference = 'Stop'; $app = New-Object -ComObject {app}; "
        f"try {{ {steps}; $doc.Close() }} finally {{ $app.Quit() }}"
    )
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True,
        text=True,
        timeout=300,
    )
    if result.returncode != 0 or not target.exists():
        raise ValueError(f"用 Office 轉換 {path.name} 失敗（需要安裝 PowerPoint／Word）：{result.stderr.strip()}")
    return target


def extract_markdown(path: Path, convert_dir: Path | None = None, image_dir: Path | None = None) -> str:
    """convert_dir：舊版 .ppt/.doc 轉出的新格式放哪裡，預設和原檔同一個資料夾。
    image_dir：投影片圖片存到哪裡（只有 PPTX 會存），不給就只標示 [圖片]。"""
    source = path
    if path.suffix.lower() in _LEGACY:
        path = convert_legacy(path, convert_dir or path.parent)
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ValueError(f"不支援 {path.suffix}，目前支援：{sorted([*_EXTRACTORS, *_LEGACY])}")
    return f"# {source.name}\n\n" + "\n\n".join(extractor(path, image_dir)) + "\n"
