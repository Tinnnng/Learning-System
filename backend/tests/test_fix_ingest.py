"""ingest 稽核（2026-10-17）找到的 bug 的回歸測試。系統內 AI 一律用 FakeProvider。"""

import io
import json
import re
import shutil

import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.modules.ai import service as ai_service
from app.modules.ai.providers.fake import FakeProvider
from app.modules.ingest import service, uploads
from app.modules.ingest.adapters.text_file import TextFormatError, parse_text_questions
from tests.test_uploads import _upload_paper, fixed_draft  # noqa: F401  （fixture）

HEAD = "# 科目：測試\n# 代碼：fix-text\n"


@pytest.fixture(autouse=True)
def clean_dirs(data_dir):
    for folder in ["inbox", "converted", "drafts", "assets"]:
        shutil.rmtree(data_dir / folder, ignore_errors=True)
    yield


@pytest.fixture
def fake_ai():
    provider = FakeProvider()
    ai_service.set_provider(provider)
    yield provider
    ai_service.set_provider(None)


@pytest.fixture
def raw_client():
    return TestClient(create_app(), raise_server_exceptions=False)


def _blank_pdf() -> bytes:
    document = pymupdf.open()
    document.new_page()
    return document.tobytes()


# ---------- 題目文字檔 ----------


def test_numbered_list_in_explanation_stays_in_explanation():
    text = HEAD + "1. 下列何者正確？\n(A) 甲\n(B) 乙\n答案：B\n詳解：理由如下\n1. 甲錯\n2. 乙對\n2. 下一題\n(A) 丙\n(B) 丁\n答案：A\n"
    bundle = parse_text_questions(text, "x")
    assert [q.number for q in bundle.questions] == [1, 2]
    assert bundle.questions[0].explanation == "理由如下\n1. 甲錯\n2. 乙對"
    assert bundle.questions[1].stem == "下一題"


def test_decimal_at_line_start_is_not_a_question():
    text = HEAD + "1. 圓周率約為\n3.14，請問下列何者正確？\n(A) 甲\n(B) 乙\n答案：A\n"
    bundle = parse_text_questions(text, "x")
    assert len(bundle.questions) == 1 and bundle.questions[0].stem == "圓周率約為\n3.14，請問下列何者正確？"


def test_code_block_keeps_indentation_and_blank_lines():
    code = "```python\ndef f():\n    if True:\n\n        return 1\n```"
    text = HEAD + f"1. 下列程式輸出為何？\n{code}\n(A) 1\n(B) 2\n答案：A\n"
    bundle = parse_text_questions(text, "x")
    assert bundle.questions[0].stem == f"下列程式輸出為何？\n{code}"
    assert bundle.questions[0].options == ["1", "2"]


def test_too_many_options_is_a_format_error():
    options = "".join(f"({c}) 選項{c}\n" for c in "ABCDEFGHIJ")
    with pytest.raises(TextFormatError, match="選項最多"):
        parse_text_questions(HEAD + "1. 題目\n" + options + "(A) 多一行\n答案：A\n", "x")


def test_fullwidth_answer_letter():
    bundle = parse_text_questions(HEAD + "1. 題目\n(A) 甲\n(B) 乙\n答案：Ｂ\n", "x")
    assert bundle.questions[0].answer == ["B"]


def test_broken_text_file_does_not_break_listing_preview_or_import(raw_client):
    service.converted_dir().mkdir(parents=True, exist_ok=True)
    options = "".join(f"({c}) 選項{c}\n" for c in "ABCDEFGHIJ")
    (service.converted_dir() / "壞檔.txt").write_text(HEAD + "1. 題目\n" + options + "(A) 多\n答案：A\n", encoding="utf-8")
    (service.converted_dir() / "沒科目.txt").write_text("1. Q\n(A) a\n(B) b\n答案：A\n", encoding="utf-8")
    listing = raw_client.get("/api/ingest/files")
    assert listing.status_code == 200
    assert all(item["error"] for item in listing.json())
    for name in ("壞檔.txt", "沒科目.txt"):
        preview = raw_client.get(f"/api/ingest/files/{name}")
        imported = raw_client.post(f"/api/ingest/files/{name}/import")
        assert preview.status_code == imported.status_code == 422
        assert "格式" in preview.json()["detail"]


# ---------- 上傳 ----------


def test_reupload_matches_earlier_numbered_copies(client):
    def up(body: bytes):
        return client.post("/api/ingest/uploads", data={"kind": "text"}, files=[("files", ("題目.txt", body, "text/plain"))]).json()["saved"]

    assert up(b"v1") == ["題目.txt"]
    assert up(b"v2") == ["題目-2.txt"]
    assert up(b"v2") == ["題目-2.txt"]
    assert up(b"v1") == ["題目.txt"]
    assert sorted(p.name for p in service.converted_dir().iterdir()) == ["題目-2.txt", "題目.txt"]


def test_corrupt_pdf_is_rejected_on_upload(client):
    response = client.post(
        "/api/ingest/uploads",
        data={"kind": "gaopu", "exam_type": "普考", "category": "一般行政"},
        files=[("files", ("115080_4401_測試科目.pdf", b"not really a pdf", "application/pdf"))],
    )
    assert response.status_code == 400 and "PDF" in response.json()["detail"]
    lecture = client.post("/api/ingest/uploads", data={"kind": "lecture"}, files=[("files", ("教材.pdf", b"broken", "application/pdf"))])
    assert lecture.status_code == 400


def test_pending_list_survives_corrupt_pdf_already_on_disk(client):
    folder = uploads.gaopu_root() / "115普考QA-一般行政"
    folder.mkdir(parents=True)
    (folder / "115080_4401_測試科目.pdf").write_bytes(b"not really a pdf")
    response = client.get("/api/ingest/uploads")
    assert response.status_code == 200
    assert response.json()["papers"][0]["kind"] == "unknown"


@pytest.mark.parametrize("category", ["資訊:處理", "a/b", "問號?", "結尾."])
def test_category_with_invalid_folder_chars(client, category):
    response = client.post(
        "/api/ingest/uploads",
        data={"kind": "gaopu", "exam_type": "普考", "category": category},
        files=[("files", ("115080_4401_測試科目.pdf", _blank_pdf(), "application/pdf"))],
    )
    assert response.status_code == 400


def test_batch_is_validated_before_anything_is_written(client):
    response = client.post(
        "/api/ingest/uploads",
        data={"kind": "gaopu", "exam_type": "普考", "category": "一般行政"},
        files=[
            ("files", ("115080_4401_測試科目.pdf", _blank_pdf(), "application/pdf")),
            ("files", ("考卷.pdf", _blank_pdf(), "application/pdf")),
        ],
    )
    assert response.status_code == 400
    assert not uploads.gaopu_root().exists() or not list(uploads.gaopu_root().rglob("*.pdf"))


def test_proofread_accepts_string_numbers(client, tmp_path, fixed_draft, fake_ai):  # noqa: F811
    _upload_paper(client, tmp_path)
    fake_ai.replies = [
        json.dumps({"questions": [{"number": "1", "stem": "$(535)_8$ 等於十進位多少？", "options": ["347", "349", "351", "353"], "note": "補上下標"}]}),
        json.dumps({"questions": [{"number": "第2題", "stem": "下列何者正確？", "options": ["甲", "乙", "丙", "丁"], "note": ""}]}),
    ]
    result = client.post("/api/ingest/uploads/papers/gaopu-115080-4401/convert", json={"use_ai": True}).json()
    bundle = service.load_bundle(service.converted_dir() / result["file"])
    assert bundle.questions[0].stem == "$(535)_8$ 等於十進位多少？"
    assert not any("沒有回傳" in note for notes in bundle.review_notes.values() for note in notes)


# ---------- 教材出題 ----------


def _scanned_pptx(path, slides: int, ext: str = "png"):
    from pptx import Presentation
    from pptx.util import Inches

    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pixmap.clear_with(200)
    data = pixmap.tobytes("png") if ext == "png" else pixmap.tobytes("jpg")
    presentation = Presentation()
    for _ in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(data), Inches(1), Inches(1))
    presentation.save(path)


def test_scanned_pptx_images_do_not_multiply(client, fake_ai):
    uploads.lecture_root().mkdir(parents=True)
    _scanned_pptx(uploads.lecture_root() / "掃描.pptx", slides=5)
    calls = [client.get("/api/ingest/uploads/lectures/掃描.pptx/estimate").json()["calls"] for _ in range(3)]
    assert calls == [1, 1, 1]
    fake_ai.replies = [json.dumps({"questions": []})]
    client.post("/api/ingest/uploads/lectures/掃描.pptx/generate", json={"count": 3})
    assert [len(call["images"]) for call in fake_ai.calls] == [5]


def test_jpg_slide_is_used_as_source_snapshot(client, fake_ai):
    uploads.lecture_root().mkdir(parents=True)
    _scanned_pptx(uploads.lecture_root() / "jpg.pptx", slides=2, ext="jpg")
    question = {"stem": "Q", "options": ["a", "b", "c", "d"], "answer": "A", "explanation": "因為 a", "source": 2}
    fake_ai.replies = [json.dumps({"questions": [question]})]
    result = client.post("/api/ingest/uploads/lectures/jpg.pptx/generate", json={"count": 1}).json()
    bundle = service.load_bundle(service.converted_dir() / result["file"])
    snapshot = bundle.questions[0].page_snapshot
    assert snapshot and snapshot.endswith((".jpg", ".jpeg")) and (service.assets_dir() / snapshot).exists()


@pytest.mark.parametrize(("count", "expected"), [(4, [2, 1, 1]), (1, [1]), (3, [1, 1, 1]), (7, [3, 2, 2])])
def test_batch_quotas_add_up_to_requested_count(client, fake_ai, monkeypatch, count, expected):
    uploads.lecture_root().mkdir(parents=True)
    (uploads.lecture_root() / "x.pdf").write_bytes(_blank_pdf())
    body = "內容" * 2600  # 每頁 5200 字，一頁一批
    markdown = "# x.pdf\n\n" + "\n\n".join(f"## 第 {n} 頁\n\n{body}" for n in (1, 2, 3))
    monkeypatch.setattr(uploads.lecture, "extract_markdown", lambda *args, **kwargs: markdown)
    fake_ai.replies = [json.dumps({"questions": []})] * 3
    client.post("/api/ingest/uploads/lectures/x.pdf/generate", json={"count": count})
    asked = [int(re.search(r"出 (\d+) 題", call["prompt"])[1]) for call in fake_ai.calls]
    assert asked == expected and sum(asked) == count


def test_docx_heading_numbers_are_not_page_numbers(client, fake_ai):
    from docx import Document

    document = Document()
    document.add_heading("3.2 迴圈", level=1)
    document.add_paragraph("for 迴圈會重複執行區塊裡的程式。" * 20)
    uploads.lecture_root().mkdir(parents=True)
    document.save(uploads.lecture_root() / "講義.docx")
    question = {"stem": "Q", "options": ["a", "b", "c", "d"], "answer": "A", "explanation": "因為 a", "source": "3.2 迴圈"}
    fake_ai.replies = [json.dumps({"questions": [question]})]
    result = client.post("/api/ingest/uploads/lectures/講義.docx/generate", json={"count": 1}).json()
    assert "Word 教材" in fake_ai.calls[0]["prompt"]
    ref = service.load_bundle(service.converted_dir() / result["file"]).questions[0].source_ref
    assert (ref.page, ref.slide, ref.note) == (None, None, "3.2 迴圈")
    assert result["report"]["ok"]
