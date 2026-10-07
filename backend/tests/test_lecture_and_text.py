from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches

from app.modules.ingest import lecture, service
from app.modules.ingest.adapters import adapter_for
from app.modules.ingest.adapters.text_file import TextFormatError, parse_text_questions
from app.modules.ingest.validation import check_bundle
from tests.factories import make_bundle

SAMPLE = Path(__file__).resolve().parents[2] / "docs" / "examples" / "sample-questions.txt"


def test_sample_text_file_parses_and_passes_validation():
    bundle = adapter_for(SAMPLE).load(SAMPLE)
    assert bundle.paper.source_key == "prog-midterm-115-sample"
    assert bundle.paper.subject == "程式設計"
    assert bundle.paper.year == 115
    assert bundle.paper.origin == "manual"
    assert [q.answer for q in bundle.questions] == [["B"], ["D"], ["A", "B", "D"]]
    assert bundle.questions[2].type == "multi_choice"
    assert "```c" in bundle.questions[1].stem
    assert bundle.questions[0].explanation.startswith("int 是")
    assert check_bundle(bundle).ok


def test_text_file_multiline_and_alternate_markers():
    content = "# 科目：測試\n\n1、第一行\n第二行\nA. 甲\nB. 乙\n答案：b\n"
    question = parse_text_questions(content, default_key="text-abc").questions[0]
    assert question.stem == "第一行\n第二行"
    assert question.options == ["甲", "乙"]
    assert question.answer == ["B"]


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("1. 題目\n(A) 甲\n答案：A\n", "缺少「# 科目"),
        ("# 科目：x\n1. 題目\n(A) 甲\n(C) 丙\n", r"應該是 \(B\)"),
        ("# 科目：x\n# 難度：高\n", "不認得的設定"),
        ("# 科目：x\n這行不是題目\n", "第一題之前"),
    ],
)
def test_text_file_errors(content, message):
    with pytest.raises(TextFormatError, match=message):
        parse_text_questions(content, default_key="text-abc")


def test_generated_questions_require_explanation_and_source():
    bundle = make_bundle(origin="generated")
    messages = [issue.message for issue in check_bundle(bundle).issues if issue.level == "error"]
    assert any("詳解" in m for m in messages)
    assert any("出處" in m for m in messages)


def test_text_files_are_listed_and_importable(client, data_dir):
    converted = data_dir / "converted"
    converted.mkdir(exist_ok=True)
    target = converted / "期中考.txt"
    target.write_text(SAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    try:
        files = {f["name"]: f for f in client.get("/api/ingest/files").json()}
        assert files["期中考.txt"]["report"]["ok"]
        outcome = client.post("/api/ingest/files/期中考.txt/import").json()
        assert outcome["result"]["questions_created"] == 3
    finally:
        target.unlink()


def test_text_files_with_same_stem_get_different_keys(tmp_path):
    from app.modules.ingest.adapters.text_file import TextFileAdapter

    body = "# 科目：程式設計\n\n1. 題目\n(A) 甲\n(B) 乙\n答案：A\n"
    (tmp_path / "期中考.txt").write_text(body, encoding="utf-8")
    (tmp_path / "期中考.md").write_text(body.replace("程式設計", "資料結構"), encoding="utf-8")
    keys = {TextFileAdapter().load(tmp_path / name).paper.source_key for name in ["期中考.txt", "期中考.md"]}
    assert len(keys) == 2


def test_extract_pptx_docx_pdf(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "指標"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "*p 取得指向的值"
    presentation.save(tmp_path / "ch3.pptx")

    document = DocxDocument()
    document.add_heading("遞迴", level=1)
    document.add_paragraph("函式呼叫自己")
    document.save(tmp_path / "notes.docx")

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Page one text")
    pdf.save(tmp_path / "slides.pdf")

    pptx_text = lecture.extract_markdown(tmp_path / "ch3.pptx")
    assert "## 投影片 1" in pptx_text and "*p 取得指向的值" in pptx_text
    docx_text = lecture.extract_markdown(tmp_path / "notes.docx")
    assert "## 遞迴" in docx_text and "函式呼叫自己" in docx_text
    assert "## 第 1 頁" in lecture.extract_markdown(tmp_path / "slides.pdf")
    with pytest.raises(ValueError, match="不支援"):
        lecture.extract_markdown(tmp_path / "slides.odp")
    with pytest.raises(ValueError, match="找不到"):
        lecture.extract_markdown(tmp_path / "old.ppt")


def test_extract_pptx_saves_slide_pictures(tmp_path):
    picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), False)
    picture.clear_with(200)
    (tmp_path / "scan.png").write_bytes(picture.tobytes("png"))
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(tmp_path / "scan.png"), Inches(1), Inches(1))
    presentation.save(tmp_path / "scanned.pptx")

    assert "[圖片]" in lecture.extract_markdown(tmp_path / "scanned.pptx")
    text = lecture.extract_markdown(tmp_path / "scanned.pptx", image_dir=tmp_path / "images")
    saved = tmp_path / "images" / "slide-01-1.png"
    assert saved.exists() and f"[圖片：{saved.as_posix()}]" in text


def test_list_converted_ignores_unsupported_files(db, data_dir):
    converted = data_dir / "converted"
    converted.mkdir(exist_ok=True)
    (converted / "readme.pdf").write_bytes(b"x")
    try:
        assert service.list_converted(db) == []
    finally:
        (converted / "readme.pdf").unlink()
